"""F4-DAS-01..03 Indicadores con formula documentada, filtros, cache, exportables y
reportes programados."""

import io
import time
from datetime import date, timedelta
from decimal import Decimal

import pytest
from httpx import AsyncClient

from app.core.enums import EstadoServicio, Rol, TipoServicio
from app.core.tiempo import ahora, hoy
from app.integrations import correo
from app.modules.cartera import service as cartera
from app.modules.dashboard import reportes, service
from app.modules.dashboard.models import ReporteProgramado
from app.modules.encuestas.models import Encuesta
from app.modules.familias import service as familias
from app.modules.rutas.models import ParadaRuta, Ruta
from app.modules.servicios.models import Servicio
from tests import fabricas
from tests.conftest import crear_usuario, login

pytestmark = pytest.mark.integration

HOY = hoy()
DESDE, HASTA = HOY - timedelta(days=10), HOY


@pytest.fixture
async def operacion(db):
    """Un mes de operacion pequeno pero completo, en Cali y Bogota."""
    cali1, _ = await fabricas.afiliada(db, HOY - timedelta(days=2), telefono="+573001000001")
    cali2, _ = await fabricas.afiliada(db, HOY - timedelta(days=1), telefono="+573001000002")
    bogota, _ = await fabricas.afiliada(
        db, HOY - timedelta(days=3), ciudad="Bogota", telefono="+573001000003"
    )
    # Una familia que paga su primera cuota y otra que se retira.
    await db.run_sync(
        lambda s: cartera.registrar_pago_manual(
            s,
            familia_id=cali1.id,
            valor=Decimal(45_000),
            medio="efectivo",
            referencia="EF-1",
            fecha_pago=None,
            cuota_id=None,
            notas=None,
            usuario_id=None,
        )
    )
    await db.run_sync(familias.dar_de_baja, bogota, "se mudo")
    # Una consulta completada, encuestada con 4, dentro de una ruta cumplida.
    servicio = Servicio(
        familia_id=cali1.id,
        tipo_codigo="consulta_domiciliaria",
        tipo=TipoServicio.DOMICILIARIO,
        ciudad="Cali",
        programado_para=ahora() - timedelta(days=1),
        estado=EstadoServicio.COMPLETADO,
    )
    db.add(servicio)
    await db.flush()
    db.add(
        Encuesta(
            servicio_id=servicio.id,
            familia_id=cali1.id,
            token="t1",
            programada_para=ahora(),
            enviada_en=ahora(),
            respondida_en=ahora(),
            puntaje=4,
            estado="respondida",
        )
    )
    ruta = Ruta(fecha=HOY - timedelta(days=1), ciudad="Cali", estado="finalizada")
    db.add(ruta)
    await db.flush()
    estimada = ahora() - timedelta(days=1)
    db.add(
        ParadaRuta(
            ruta_id=ruta.id,
            servicio_id=servicio.id,
            orden=1,
            estado="completada",
            llegada_estimada=estimada,
            llegada_en=estimada + timedelta(minutes=10),
        )
    )
    await db.flush()


async def _calcular(db, ciudad=None, codigos=None):
    return (await db.run_sync(service.calcular, DESDE, HASTA, ciudad, codigos))["indicadores"]


# --- F4-DAS-01 / 02 Indicadores ---------------------------------------------------------------


async def test_indicadores_del_periodo(db, operacion):
    i = await _calcular(db)
    assert (i["afiliaciones"], i["retiros"]) == (3, 1)
    assert i["recaudo"] == 45_000.0
    assert i["facturado"] == 3 * 45_000.0
    assert i["servicios_por_tipo"] == {"consulta_domiciliaria": 1}
    assert i["satisfaccion"]["promedio"] == 4.0 and i["satisfaccion"]["tasa_respuesta"] == 1.0
    assert i["cumplimiento_rutas"]["cumplimiento"] == 1.0
    assert i["cumplimiento_rutas"]["puntualidad"] == 1.0


async def test_filtro_por_ciudad(db, operacion):
    i = await _calcular(db, ciudad="Cali", codigos=["afiliaciones", "retiros", "recaudo"])
    assert i == {"afiliaciones": 2, "retiros": 0, "recaudo": 45_000.0}


async def test_mora_al_cierre(db):
    await fabricas.afiliada(db, HOY - timedelta(days=40), telefono="+573001000009")
    await db.run_sync(cartera.generar_cuotas, HOY)
    i = await _calcular(db, codigos=["mora"])
    assert i["mora"]["contratos_en_mora"] == 1 and i["mora"]["saldo_vencido"] > 0


async def test_cache_y_rendimiento(db, operacion):
    inicio = time.perf_counter()
    primera = await _calcular(db)
    assert time.perf_counter() - inicio < 2  # F4-DAS-02 menos de dos segundos
    await fabricas.afiliada(db, HOY, telefono="+573001000010")
    assert (await _calcular(db))["afiliaciones"] == primera["afiliaciones"]  # cacheado
    from app.core import cache

    cache.obtener().limpiar("dashboard:")
    assert (await _calcular(db))["afiliaciones"] == primera["afiliaciones"] + 1


async def test_periodo_invalido(client: AsyncClient, headers):
    resp = await client.get(
        "/api/v1/dashboard/indicadores",
        params={"desde": str(HASTA), "hasta": str(DESDE)},
        headers=headers,
    )
    assert resp.status_code == 422


async def test_definiciones_documentadas(client: AsyncClient, headers):
    defs = (await client.get("/api/v1/dashboard/definiciones", headers=headers)).json()
    assert {d["codigo"] for d in defs} == set(service.INDICADORES)
    assert all(len(d["formula"]) > 20 for d in defs)


async def test_el_medico_no_ve_el_dashboard(client: AsyncClient, db):
    await crear_usuario(db, "vet@timu.co", Rol.MEDICO)
    resp = await client.get(
        "/api/v1/dashboard/indicadores",
        params={"desde": str(DESDE), "hasta": str(HASTA)},
        headers=await login(client, "vet@timu.co"),
    )
    assert resp.status_code == 403


# --- F4-DAS-03 Exportables y reportes ---------------------------------------------------------


async def test_exporta_csv_y_excel(client: AsyncClient, db, operacion, headers):
    await db.commit()
    params = {"desde": str(DESDE), "hasta": str(HASTA)}
    csv = await client.get(
        "/api/v1/dashboard/indicadores", params={**params, "formato": "csv"}, headers=headers
    )
    assert csv.headers["content-type"].startswith("text/csv")
    assert "Afiliaciones" in csv.content.decode("utf-8-sig")

    xlsx = await client.get(
        "/api/v1/dashboard/indicadores", params={**params, "formato": "xlsx"}, headers=headers
    )
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(xlsx.content))
    assert libro.sheetnames == ["Indicadores", "Definiciones"]
    valores = [fila[0] for fila in libro["Indicadores"].iter_rows(values_only=True)]
    assert "Afiliaciones" in valores


def test_periodo_anterior():
    lunes = date(2026, 10, 19)
    assert reportes.periodo_anterior("semanal", lunes) == (date(2026, 10, 12), date(2026, 10, 18))
    assert reportes.periodo_anterior("mensual", date(2026, 11, 1)) == (
        date(2026, 10, 1),
        date(2026, 10, 31),
    )
    assert reportes.periodo_anterior("diaria", lunes) == (date(2026, 10, 18), date(2026, 10, 18))


async def test_reporte_semanal_por_correo(client: AsyncClient, db, operacion, headers, monkeypatch):
    enviados = []
    monkeypatch.setattr(
        correo,
        "enviar",
        lambda destino, asunto, cuerpo, adjuntos=None: enviados.append((destino, asunto, adjuntos))
        or "enviado",
    )
    await db.commit()
    malo = await client.post(
        "/api/v1/dashboard/reportes",
        json={
            "nombre": "x",
            "periodicidad": "semanal",
            "destinatarios": ["a@b.co"],
            "indicadores": ["inventado"],
        },
        headers=headers,
    )
    assert malo.status_code == 422
    creado = await client.post(
        "/api/v1/dashboard/reportes",
        json={
            "nombre": "Gerencia semanal",
            "periodicidad": "semanal",
            "destinatarios": ["gerencia@timu.co", "socios@timu.co"],
        },
        headers=headers,
    )
    assert creado.status_code == 201, creado.text

    lunes = HOY - timedelta(days=HOY.weekday())
    martes = lunes + timedelta(days=1)
    assert (await db.run_sync(reportes.enviar_programados, martes))["enviados"] == 0
    assert (await db.run_sync(reportes.enviar_programados, lunes))["enviados"] == 1
    assert [d for d, _, _ in enviados] == ["gerencia@timu.co", "socios@timu.co"]
    nombre, contenido, tipo = enviados[0][2][0]
    assert nombre.endswith(".xlsx") and contenido.startswith(b"PK")
    # El mismo dia no sale dos veces.
    reporte = await db.get(ReporteProgramado, creado.json()["id"])
    reporte.ultimo_envio = ahora()
    assert (await db.run_sync(reportes.enviar_programados, HOY))["enviados"] == 0
